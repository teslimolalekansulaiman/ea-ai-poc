"""Export eval/gold.json to a shareable Excel workbook (eval/gold.xlsx).

Sheets: README (the note + doc scope), Entities, Relations, Review (items to verify).
For the EA team to review/extend the gold set in a familiar tool.

    python eval/gold_to_xlsx.py            # -> eval/gold.xlsx
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
HEAD_FILL = PatternFill("solid", fgColor="2D6CC0")
HEAD_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")


def _sheet(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c, w in zip(ws[1], widths):
        c.fill = HEAD_FILL; c.font = HEAD_FONT
        ws.column_dimensions[c.column_letter].width = w
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    ws.freeze_panes = "A2"
    return ws


def main() -> None:
    gold = json.loads((ROOT / "eval/gold.json").read_text(encoding="utf-8"))
    wb = Workbook()
    wb.remove(wb.active)

    # README
    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 110
    ws["A1"] = "EA Knowledge-Graph — Gold Evaluation Set"
    ws["A1"].font = Font(bold=True, size=14)
    lines = [
        "",
        "Purpose: the hand-labelled 'answer key' used to measure the extraction pipeline "
        "(precision / recall / F1). Please REVIEW and EXTEND it against the original .docx files.",
        "",
        f"Documents in scope: {', '.join(gold['docs'])}",
        f"Entities: {len(gold['entities'])}   Relations: {len(gold['relations'])}   "
        f"Review items: {len(gold.get('review', []))}",
        "",
        gold.get("_note", ""),
        "",
        "How to help:",
        "  • Entities / Relations tabs: add anything the documents state that's missing; fix any that are wrong.",
        "  • Review tab: these are things the graph produced that look like errors/ambiguities — confirm or correct.",
        "  • Send the edited file back and we fold it into eval/gold.json.",
    ]
    for i, ln in enumerate(lines, start=2):
        ws[f"A{i}"] = ln
        ws[f"A{i}"].alignment = WRAP

    _sheet(wb, "Entities", ["Entity name", "Type"],
           [[e["name"], e["type"]] for e in gold["entities"]], [45, 32])
    _sheet(wb, "Relations", ["Source", "Relation", "Target"],
           [[r["source"], r["relation"], r["target"]] for r in gold["relations"]], [34, 22, 44])
    _sheet(wb, "Review (verify these)", ["Item", "Graph type(s)", "Note / question"],
           [[r.get("item", ""), r.get("graph_type", ""), r.get("note", "")]
            for r in gold.get("review", [])], [32, 30, 80])

    out = ROOT / "eval/gold.xlsx"
    wb.save(out)
    print(f"wrote {out}  ({len(gold['entities'])} entities, {len(gold['relations'])} relations, "
          f"{len(gold.get('review', []))} review items)")


if __name__ == "__main__":
    main()
