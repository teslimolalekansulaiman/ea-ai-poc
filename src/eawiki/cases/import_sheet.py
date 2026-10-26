"""Load the EA team's verified guidance sheet into a use case.

The sheet has one row per property: its description, the steps, where to look, the rules, the
examples, the pick list, what it builds on, whether a source is required and how many, and what to
ask when input or a source is missing. This reads it, matches each row to a property of the use case
by name, and reports what it would change. Nothing is written unless asked, and what is written goes
through the same validation as the Config screen.

    python -m eawiki.cases.import_sheet <sheet.xlsx> [--use-case demand-management] [--write] [--include-drafts]

Rules for what is read:
  * a note in brackets for the reader ("[PROPOSED by Claude ...]", "[EA to write]") is removed; a
    field that is only such a note is left alone;
  * a row that matches no property is reported, never added;
  * a sheet marked 'draft' (properties drafted for EA to write) is skipped unless --include-drafts;
  * the row about the stage ("Stage instructions") becomes the use case's rules for every property.
"""

from __future__ import annotations

import re
import sys
from typing import Optional

from eawiki import usecases

NOTE = re.compile(r"\[(?:PROPOSED|EA to)[^\]]*\]\s*", re.I)
TEXT_FIELDS = {"description": "description", "how to resolve": "how_to_resolve", "where to look": "where_to_look",
               "rules": "rules", "examples": "examples"}


def _words(s: str) -> set:
    return set(re.findall(r"[^\W_]{2,}", str(s or "").casefold()))


def match_column(field: str, names: list[str]) -> Optional[str]:
    norm = lambda x: re.sub(r"\s+", " ", x).strip().casefold()
    exact = [n for n in names if norm(n) == norm(field)]
    if len(exact) == 1:
        return exact[0]
    sub = [n for n in names if _words(field) and _words(field) <= _words(n)]
    return sub[0] if len(sub) == 1 else None


def _clean(v) -> str:
    return NOTE.sub("", str(v or "")).strip()


def _header_map(header: list) -> dict:
    out = {}
    for i, h in enumerate(header):
        t = str(h or "").strip().casefold()
        if t == "field":
            out["field"] = i
        elif t.startswith("values"):
            out["values"] = i
        elif t.startswith("builds on"):
            out["builds"] = i
        elif t.startswith("source required"):
            out["src"] = i
        elif t.startswith("min sources"):
            out["min"] = i
        elif t.startswith("info needed") and "input" in t:
            out["info_input"] = i
        elif t.startswith("info needed") and "source" in t:
            out["info_source"] = i
        else:
            for key in TEXT_FIELDS:
                if t.startswith(key):
                    out[key] = i
    return out


def read_rows(path: str, drafts: bool = False) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    rows = []
    for ws in wb.worksheets:
        if "draft" in ws.title.casefold() and not drafts:
            continue                      # properties Claude drafted for EA to write: never loaded as verified
        data = list(ws.iter_rows(values_only=True))
        if not data:
            continue
        hm = _header_map(list(data[0]))
        if "field" not in hm:
            continue
        for r in data[1:]:
            f = r[hm["field"]] if hm["field"] < len(r) else None
            if f and str(f).strip() and len(str(f)) <= 80:
                rows.append({"sheet": ws.title, "field": str(f).strip(),
                             **{k: (r[i] if i < len(r) else None) for k, i in hm.items() if k != "field"}})
    return rows


def apply(spec: dict, path: str, drafts: bool = False) -> tuple[dict, dict]:
    """(new spec, report). The given spec is not changed."""
    import json
    spec = json.loads(json.dumps(spec))
    cols = {c["name"]: c for st in spec["stages"] for c in st["output"]["columns"]}
    names = list(cols)
    rep = {"changed": {}, "unmatched": [], "skipped": [], "playbook": False}
    for row in read_rows(path, drafts):
        field = row["field"]
        if field.casefold().startswith("stage instructions"):
            rules = _clean(row.get("rules"))
            if rules:
                spec["playbook"] = rules
                rep["playbook"] = True
            continue
        name = match_column(field, names)
        if not name:
            rep["unmatched"].append(field)
            continue
        c, did = cols[name], []
        for key, attr in TEXT_FIELDS.items():
            v = _clean(row.get(key))
            if v and v != (c.get(attr) or ""):
                c[attr] = v
                did.append(attr)
        for key, attr in (("info_input", "info_needed"), ("info_source", "info_needed_source")):
            v = _clean(row.get(key))
            if v and v != (c.get(attr) or ""):
                c[attr] = v
                did.append(attr)
        vals = [x.strip() for x in re.split(r"[;\n]", _clean(row.get("values"))) if x.strip()]
        if vals and vals != (c.get("values") or []):
            c["values"] = vals
            did.append("values")
        uses = [match_column(x.strip(), [n for n in names if n != name]) for x in re.split(r"[;\n]", _clean(row.get("builds")))
                if x.strip() and not x.strip().casefold().startswith("all the")]
        uses = list(dict.fromkeys(u for u in uses if u))
        if uses and uses != (c.get("uses") or []):
            c["uses"] = uses
            did.append("uses")
        src = _clean(row.get("src")).casefold()
        if src.startswith("no") and c.get("evidence_required") is not False:
            c["evidence_required"] = False
            did.append("evidence_required")
        elif src.startswith("yes") and c.get("evidence_required") is False:
            c.pop("evidence_required")
            did.append("evidence_required")
        m = re.match(r"\s*(\d)", _clean(row.get("min")))
        if m and c.get("min_sources") != int(m.group(1)):
            c["min_sources"] = int(m.group(1))
            did.append("min_sources")
        if did:
            rep["changed"][name] = did
        else:
            rep["skipped"].append(name)
    return spec, rep


def main(argv: list[str]) -> None:
    if not argv:
        print(__doc__)
        return
    uc = argv[argv.index("--use-case") + 1] if "--use-case" in argv else "demand-management"
    spec = usecases.get(uc)
    if not spec:
        sys.exit(f"no use case {uc}")
    new, rep = apply(spec, argv[0], "--include-drafts" in argv)
    for n, d in rep["changed"].items():
        print(f"  {n}: {', '.join(d)}")
    if rep["playbook"]:
        print("  use case rules (playbook): set")
    if rep["unmatched"]:
        print("  not in the use case (add the column first):", ", ".join(rep["unmatched"]))
    if rep["skipped"]:
        print("  nothing to change:", ", ".join(rep["skipped"]))
    errs = usecases.validate(new)
    if errs:
        sys.exit("would not validate:\n  " + "\n  ".join(errs))
    if "--write" in argv:
        usecases.save(new, uc)
        print("written.")
    else:
        print("dry run: nothing written. Add --write to save.")


if __name__ == "__main__":
    main(sys.argv[1:])
