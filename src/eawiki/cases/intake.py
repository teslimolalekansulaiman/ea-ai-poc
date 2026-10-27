"""Read a sheet of items (CSV or Excel) into the use case's input fields.

Headers are matched to fields by the field's label or name, ignoring case and a trailing '*'. A
row that only repeats the guidance text (as the EA's sheet does under its headers) is skipped. A
header that matches no field is reported, not silently dropped.
"""

from __future__ import annotations

import csv
import io
import re
from typing import Optional


def _norm(s) -> str:
    return re.sub(r"[\s*]+", " ", str(s or "")).strip().lower().rstrip("*").strip()


def field_index(spec: dict) -> dict:
    idx = {}
    for f in (spec.get("input") or {}).get("item_fields") or []:
        idx[_norm(f["name"])] = f["name"]
        idx[_norm(f.get("label") or f["name"])] = f["name"]
    return idx


def rows_to_items(spec: dict, rows: list[list]) -> dict:
    idx = field_index(spec)
    guidance = {f["name"]: _norm(f.get("description")) for f in (spec.get("input") or {}).get("item_fields") or []}
    head = None
    for i, r in enumerate(rows[:10]):
        if sum(1 for c in r if _norm(c) in idx) >= 3:
            head = i
            break
    if head is None:
        return {"items": [], "unmatched": [], "error": "No header row found: the sheet needs the column names of the input fields."}
    cols = [idx.get(_norm(c)) for c in rows[head]]
    unmatched = [str(c) for c, k in zip(rows[head], cols) if c and k is None]
    items = []
    for r in rows[head + 1:]:
        vals = {k: ("" if v is None else str(v).strip()) for k, v in zip(cols, r) if k}
        if not any(vals.values()):
            continue
        if all(_norm(v) == guidance.get(k, "") or not v for k, v in vals.items()):
            continue                                   # the guidance row under the headers
        items.append(vals)
    return {"items": items, "unmatched": unmatched, "error": ""}


def parse_file(spec: dict, name: str, data: bytes) -> dict:
    lower = (name or "").lower()
    if lower.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError:
            return {"items": [], "unmatched": [], "error": "Excel files need the openpyxl package on the server."}
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        text = data.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
    return rows_to_items(spec, rows)


_LABELLED = re.compile(r"^\s*(?:[-*\u2022]\s*)?([^:\uff1a\n]{2,60}?)\s*[:\uff1a]\s*(.*)$")
MIN_LABELLED = 3           # three recognised "Label: value" lines make a pasted demand a structured one


def _labelled(spec: dict, text: str) -> dict:
    """The "Label: value" lines of a pasted demand whose label is an input field of the use case (by its name or its label).
    A line that is not labelled continues the value above it; text before the first label is kept apart ('')."""
    idx = field_index(spec)
    out: dict = {}
    cur = ""
    for ln in text.splitlines():
        m = _LABELLED.match(ln)
        field = idx.get(_norm(m.group(1))) if m else None
        if field:
            cur = field
            out[cur] = (out[cur] + "\n" if cur in out else "") + m.group(2).strip()
        elif ln.strip():
            out[cur] = (out.get(cur, "") + ("\n" if out.get(cur) else "") + ln.strip())
    return out


def from_text(spec: dict, text: str) -> dict:
    """One typed or pasted demand. If it lists the input fields ("Department: Finance", "Business need: ..."), each line goes
    to its field and the demand is read as a sheet row would be, except that no required field is demanded: the person listed
    what they have. Otherwise the whole text is the business need and the title its first words."""
    inp = spec.get("input") or {}
    need = inp.get("batch_text_field") or "business_need"
    title_field = inp.get("title_field") or "title"
    parsed = _labelled(spec, text)
    fields = {k: v for k, v in parsed.items() if k}
    if len(fields) >= MIN_LABELLED:
        loose = parsed.get("", "").strip()
        if loose and not fields.get(need):
            fields[need] = loose                           # text outside any label is the need when none was labelled
        elif loose:
            fields[need] = (loose + "\n" + fields[need]).strip()
        if not fields.get(title_field) and (fields.get(need) or "").strip():
            fields[title_field] = " ".join(fields[need].split()[:12])
        fields["_pasted"] = True
        return {"items": [fields], "unmatched": [], "error": ""}
    title = " ".join(text.split()[:12])
    item = {need: text.strip()}
    if title_field != need:
        item[title_field] = title
    return {"items": [item], "unmatched": [], "error": ""}
