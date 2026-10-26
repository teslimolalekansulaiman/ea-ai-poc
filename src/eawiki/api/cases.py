"""HTTP routes for cases: create from a sheet or a typed demand, follow a run, answer, approve, ask.

A run executes in a background thread and writes every cell as it is decided, so the screen polls
`GET /api/cases/{id}` and sees it fill in. (On Cloud Run the service that runs cases needs CPU
allocated outside requests, as the ingest does.)
"""

from __future__ import annotations

from typing import Optional

from fastapi import APIRouter, File, HTTPException, Response, UploadFile
from pydantic import BaseModel

from eawiki import usecases
from eawiki.cases import intake, summary
from eawiki.cases.engine import Engine, gate_column, phase_id, steps
from eawiki.cases.store import store

router = APIRouter(prefix="/api")
_engine: Optional[Engine] = None


def get_engine() -> Engine:
    global _engine
    if _engine is None:
        from eawiki.cases.tools import real_tools
        _engine = Engine(store(), real_tools)
    return _engine


def set_engine(e: Optional[Engine]) -> None:
    global _engine
    _engine = e


def _spec(uc_id: str) -> dict:
    spec = usecases.get(uc_id)
    if not spec:
        raise HTTPException(404, f"no use case '{uc_id}'")
    return spec


class ParseBody(BaseModel):
    use_case_id: str
    text: str


@router.post("/cases/parse-file")
async def parse_file(use_case_id: str, file: UploadFile = File(...)):
    return intake.parse_file(_spec(use_case_id), file.filename or "", await file.read())


@router.post("/cases/parse-text")
def parse_text(body: ParseBody):
    return intake.from_text(_spec(body.use_case_id), body.text)


class CreateBody(BaseModel):
    use_case_id: str
    items: list[dict]
    name: str = ""


@router.post("/cases")
def create_case(body: CreateBody):
    if not body.items:
        raise HTTPException(400, "No demands to run.")
    spec = _spec(body.use_case_id)
    e = get_engine()
    cid = e.create(spec, body.items, body.name or f"{len(body.items)} demand(s)")
    e.start(cid)
    return {"id": cid}


@router.get("/cases")
def list_cases():
    return {"cases": store().list_cases()}


@router.delete("/cases/{cid}")
def delete_case(cid: str):
    return {"deleted": store().delete_case(cid)}


def _column_view(c: dict) -> dict:
    return {k: c.get(k) for k in ("name", "type", "values", "description", "how_to_resolve", "look_in",
                                  "evidence_required", "info_needed", "uses") if c.get(k) is not None}


@router.get("/cases/documents")
def documents():
    """The ingested documents, for the 'look here' picker."""
    return {"documents": sorted(get_engine().tools_factory().doc_ids() or [])}


@router.get("/cases/{cid}")
def get_case(cid: str, after: int = 0):
    st = store()
    case = st.get_case(cid)
    if not case:
        raise HTTPException(404, "no such case")
    spec = case["spec"]
    cells = st.cells(cid)
    questions = st.questions(cid)
    tree: dict = {}
    progress: dict = {}
    for c in cells:
        c = dict(c)
        c["display"] = summary.display_status(spec, c)
        tree.setdefault(c["phase_id"], {}).setdefault(c["item_id"], {})[c["column"]] = c
        p = progress.setdefault(c["phase_id"], {}).setdefault(c["column"], 0)
        progress[c["phase_id"]][c["column"]] = p + 1
    phases = []
    for si, stage in enumerate(spec["stages"]):
        pid = phase_id(stage)
        phases.append({"id": pid, "name": stage.get("name"), "instructions": stage.get("instructions", ""),
                       "columns": [_column_view(c) for c in stage["output"]["columns"]],
                       "state": case["state"].get("phases", {}).get(pid, {"status": "waiting"}),
                       "progress": progress.get(pid, {}), "steps": steps(spec, stage, si == 0),
                       "summary": summary.phase_summary(spec, stage, cells, case["items"], questions)})
    events = st.events_after(cid, after)
    return {"case": {"id": cid, "name": case["name"], "use_case_id": case["use_case_id"],
                     "use_case_name": spec.get("name"), "created_at": case["created_at"]},
            "input_fields": (spec.get("input") or {}).get("item_fields") or [],
            "gate_column": gate_column(spec), "items": case["items"], "phases": phases, "cells": tree,
            "questions": questions, "events": events,
            "last_seq": events[-1]["seq"] if events else after}


class AnswerBody(BaseModel):
    answer: Optional[str] = None
    mode: str = "value"                 # "value": this is the answer; "look": search where I point
    document: Optional[str] = None      # look: an ingested document (checked against the list)
    where: Optional[str] = None         # look: where in it, in words


@router.post("/cases/{cid}/questions/{qid}/answer")
def answer(cid: str, qid: str, body: AnswerBody):
    q = store().get_question(qid)
    if not q or q["case_id"] != cid:
        raise HTTPException(404, "no such question")
    if body.mode not in ("value", "look"):
        raise HTTPException(400, "mode must be 'value' or 'look'")
    try:
        if body.mode == "look":
            return get_engine().look(qid, body.document, body.where)
        return get_engine().answer(qid, body.answer)
    except ValueError as e:
        raise HTTPException(409, str(e))


@router.post("/cases/{cid}/phases/{pid}/approve")
def approve(cid: str, pid: str):
    case = store().get_case(cid)
    if not case:
        raise HTTPException(404, "no such case")
    if case["state"].get("phases", {}).get(pid, {}).get("status") != "review":
        raise HTTPException(409, "This phase is not waiting for review.")
    return get_engine().approve(cid, pid)


@router.post("/cases/{cid}/cancel")
def cancel(cid: str):
    get_engine().cancel(cid)
    return {"cancelled": True}


class CaseAsk(BaseModel):
    question: str
    history: list = []
    focus: Optional[str] = None         # the demand the architect has open, if any


def case_context(case: dict, cells: list[dict], limit: int = 14000) -> str:
    """The case's results as plain text: one block per item, every property decided so far."""
    spec = case["spec"]
    by: dict = {}
    for c in cells:
        by.setdefault((c["phase_id"], c["item_id"]), {})[c["column"]] = c
    lines = []
    for stage in spec["stages"]:
        pid = phase_id(stage)
        lines.append(f"## {stage.get('name')}")
        for it in case["items"]:
            cs = by.get((pid, it["_id"]))
            if not cs:
                continue
            title = it.get((spec.get("input") or {}).get("title_field") or "title") or ""
            lines.append(f"### {it['_id']} {title}")
            for col, cell in cs.items():
                v = "; ".join(map(str, cell["value"])) if isinstance(cell["value"], list) else cell["value"]
                lines.append(f"- {col}: {v or '(' + cell['status'] + ')'}" + (f"  [{cell['need']}]" if cell.get("need") else ""))
    return "\n".join(lines)[:limit]


@router.post("/cases/{cid}/ask")
def ask_case(cid: str, body: CaseAsk):
    case = store().get_case(cid)
    if not case:
        raise HTTPException(404, "no such case")
    from eawiki.api import app as appmod
    from eawiki.query.ask import ask as run
    ctx = case_context(case, store().cells(cid))
    if body.focus and any(i["_id"] == body.focus for i in case["items"]):
        ctx = f"The architect has {body.focus} open: questions without a demand number are about it.\n\n" + ctx
    out = run(body.question, k=8, backend=appmod.API_BACKEND, model=appmod.ASK_MODEL, use_llm=True,
              safety_llm=False, history=body.history, case_context=ctx)
    return {k: v for k, v in out.items() if k not in ("context",)}


class ReadyBody(BaseModel):
    items: list[dict] = []
    case_id: Optional[str] = None


@router.post("/usecases/{uc_id}/readiness")
def readiness(uc_id: str, body: ReadyBody):
    from eawiki.cases import readiness as rd
    spec = _spec(uc_id)
    items = body.items
    if body.case_id:
        c = store().get_case(body.case_id)
        items = c["items"] if c else items
    if not items:
        raise HTTPException(400, "Give sample demands, or a case to take them from.")
    e = get_engine()
    for i, it in enumerate(items):
        it.setdefault("_id", it.get("demand_id") or f"item-{i + 1}")
    return rd.readiness(spec, items, e.tools_factory(), e._plans)


def _show(v) -> str:
    return "; ".join(map(str, v)) if isinstance(v, list) else ("" if v is None else str(v))


@router.get("/cases/{cid}/export")
def export(cid: str, phase: str = ""):
    """The case's table for one phase (default: the first that has results) as an Excel file: one row
    per demand, one column per property, in the order the use case declares them."""
    import io
    case = store().get_case(cid)
    if not case:
        raise HTTPException(404, "no such case")
    spec = case["spec"]
    cells = store().cells(cid)
    stage = next((s for s in spec["stages"] if phase_id(s) == phase), None) or next(
        (s for s in spec["stages"] if any(c["phase_id"] == phase_id(s) for c in cells)), spec["stages"][0])
    pid = phase_id(stage)
    cols = [c["name"] for c in stage["output"]["columns"]]
    by = {(c["item_id"], c["column"]): c for c in cells if c["phase_id"] == pid}
    title = (spec.get("input") or {}).get("title_field") or "title"
    head = ["Demand", "Title"] + cols
    rows = [[it["_id"], it.get(title, "")] + [_show((by.get((it["_id"], c)) or {}).get("value")) for c in cols]
            for it in case["items"] if any((it["_id"], c) in by for c in cols)]
    name = "".join(ch if ch.isalnum() or ch in " -_" else "_" for ch in f"{case['name']} - {stage.get('name')}")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError:
        import csv
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(head)
        w.writerows(rows)
        return Response("\ufeff" + buf.getvalue(), media_type="text/csv; charset=utf-8",
                        headers={"Content-Disposition": f'attachment; filename="{name}.csv"'})
    wb = Workbook()
    ws = wb.active
    ws.title = (stage.get("name") or "Case")[:31]
    ws.append(head)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, h in enumerate(head, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 44 if h in ("Title", "Screening observation", "Justification") else 22
    ws.freeze_panes = "C2"
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}.xlsx"'})
