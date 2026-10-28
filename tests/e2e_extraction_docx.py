"""Build a Word file that is hard to extract on purpose, with a ground truth for every piece of it.

Tables (small, key-value with merged cells, a 60-row table with a repeated header, a table with a merged header, a nested table), pictures
(a capability heat map, an architecture diagram with labelled arrows, a bar chart), captions that follow their picture as a separate
paragraph (as in the real documents), lists, and a text box. The ground truth (what a correct extraction and a correct description must
contain) is written next to the file as truth.json. Opt-in script; it needs python-docx and pillow.

    python tests/e2e_extraction_docx.py <output folder>
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont

SHADES = {1: (222, 230, 242), 2: (180, 199, 225), 3: (125, 163, 207), 4: (62, 108, 160), 5: (28, 55, 92)}


def font(size, bold=False):
    try:
        return ImageFont.truetype("DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf", size)
    except Exception:
        return ImageFont.load_default(size)


CAP_MAP = {
    "Operational Excellence": [("Services and Products", 2), ("Business Procedures", 2), ("Beneficiary Experience", 2)],
    "Digital Transformation": [("Planning and Governance", 3), ("Enterprise Architecture", 2), ("Consulting Support", 2), ("Digital Projects", 3)],
    "Data Management and Governance": [("Data Management", 2), ("Data Governance", 2), ("Business Analysis", 2)],
    "Information Technology": [("Digital Applications and Platforms", 3), ("Infrastructure", 3), ("Technical and Technological Analysis", 2), ("Technological Quality", 2)],
    "Cybersecurity": [("Cybersecurity", 3)],
    "Research and Studies": [("Study and Analysis", 2), ("Innovation and Development", 1)],
}


def heat_map(path):
    W, H = 1180, 760
    im = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(im)
    cols = [["Operational Excellence", "Information Technology"], ["Digital Transformation", "Cybersecurity"], ["Data Management and Governance", "Research and Studies"]]
    for ci, groups in enumerate(cols):
        x, y = 40 + ci * 370, 40
        for g in groups:
            caps = CAP_MAP[g]
            h = 50 + len(caps) * 62
            d.rounded_rectangle([x, y, x + 340, y + h], 12, outline=(190, 200, 215), width=2, fill=(245, 248, 252))
            d.rounded_rectangle([x, y, x + 340, y + 46], 12, fill=(64, 108, 160))
            d.text((x + 170, y + 23), g, fill="white", font=font(19, True), anchor="mm")
            for i, (c, m) in enumerate(caps):
                by = y + 58 + i * 62
                d.rounded_rectangle([x + 12, by, x + 328, by + 50], 8, fill=SHADES[m])
                d.text((x + 170, by + 25), c, fill=(20, 30, 60) if m < 4 else "white", font=font(17), anchor="mm")
            y += h + 24
    lx, ly = 60, 690
    d.text((lx, ly), "Maturity:", fill=(60, 70, 90), font=font(17, True), anchor="lm")
    for i, (m, label) in enumerate([(1, "1 Ad hoc"), (2, "2 Documented"), (3, "3 Evidenced"), (4, "4 Measured"), (5, "5 Optimized")]):
        x = 170 + i * 200
        d.rectangle([x, ly - 14, x + 40, ly + 14], fill=SHADES[m])
        d.text((x + 50, ly), label, fill=(60, 70, 90), font=font(17), anchor="lm")
    im.save(path)


ARCH = {"boxes": ["HR & Payroll Portal", "Evaluate Service", "Payroll Database", "Identity Provider", "Reporting Hub"],
        "arrows": [("Evaluate Service", "HR & Payroll Portal", "REST API, daily"), ("HR & Payroll Portal", "Payroll Database", "SQL batch"),
                   ("Identity Provider", "HR & Payroll Portal", "SAML login"), ("HR & Payroll Portal", "Reporting Hub", "SFTP, nightly")]}


def architecture(path):
    im = Image.new("RGB", (1000, 560), "white")
    d = ImageDraw.Draw(im)
    pos = {"Evaluate Service": (60, 60), "HR & Payroll Portal": (400, 220), "Payroll Database": (740, 60), "Identity Provider": (60, 400), "Reporting Hub": (740, 400)}
    for n, (x, y) in pos.items():
        d.rounded_rectangle([x, y, x + 200, y + 90], 10, fill=(214, 228, 245), outline=(64, 108, 160), width=3)
        d.text((x + 100, y + 45), n, fill=(20, 30, 60), font=font(17, True), anchor="mm")
    for a, b, label in ARCH["arrows"]:
        (x1, y1), (x2, y2) = pos[a], pos[b]
        p1, p2 = (x1 + 100, y1 + 45), (x2 + 100, y2 + 45)
        d.line([p1, p2], fill=(80, 80, 80), width=4)
        mx, my = (p1[0] + p2[0]) // 2, (p1[1] + p2[1]) // 2
        d.rectangle([mx - 80, my - 14, mx + 80, my + 14], fill="white")
        d.text((mx, my), label, fill=(150, 30, 30), font=font(16), anchor="mm")
        ang = (p2[0] - p1[0], p2[1] - p1[1])
        d.ellipse([p2[0] - 9 - ang[0] * 0 // 1, p2[1] - 9, p2[0] + 9, p2[1] + 9], fill=(80, 80, 80))
    im.save(path)


BARS = [("Human Capital", 2.4), ("Finance", 2.9), ("Operations", 2.1), ("IT", 2.7), ("Legal", 1.8), ("Procurement", 2.2)]


def chart(path):
    im = Image.new("RGB", (900, 520), "white")
    d = ImageDraw.Draw(im)
    d.text((450, 30), "Average capability maturity by level-1 capability", fill=(30, 40, 70), font=font(22, True), anchor="mm")
    for i, (n, v) in enumerate(BARS):
        y = 80 + i * 70
        d.text((200, y + 22), n, fill=(30, 40, 70), font=font(18), anchor="rm")
        d.rectangle([215, y, 215 + int(v * 200), y + 44], fill=(64, 108, 160))
        d.text((215 + int(v * 200) + 10, y + 22), f"{v}", fill=(30, 40, 70), font=font(18, True), anchor="lm")
    im.save(path)


APP_ROWS = [("Description", "A service that helps the HR Department create HR evaluation forms. It is implemented inside the HR & Payroll Portal."),
            ("Responsible operator", "Application & Development Dept"), ("Application users", "Organization employees (HR Department)"),
            ("Stakeholders", "Product Manager HR Department"), ("Team Members", "Development team"), ("Application Interfaces", "Internal: Web App."),
            ("Hosting", "Operated by the hosting provider."), ("Technologies Used", "Frontend: PHP DB: MySQL"), ("Type", "Application"),
            ("Application Status", "Not in use"), ("Notes", "Replaced by the HR & Payroll Portal; used once a year for historical data only")]
OBJECTIVES = [("SO1", "Raise sector compliance", "Make compliance measurable across the sector"),
              ("SO2", "Deliver seamless digital services", "Move citizen services to online channels"),
              ("SO3", "Build national data capability", "Share trusted data across entities"),
              ("SO4", "Strengthen institutional excellence", "Streamline internal services and cut manual work in HR and finance"),
              ("SO5", "Grow the knowledge economy", "Invest in research and innovation"),
              ("SO6", "Protect digital assets", "Raise cybersecurity maturity")]


def shade_header(cell):
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="D9E2F3"/>'))


def build(out: Path):
    out.mkdir(parents=True, exist_ok=True)
    heat_map(out / "heat.png")
    architecture(out / "arch.png")
    chart(out / "chart.png")
    doc = Document()
    doc.add_heading("Test Architecture Pack", 0)
    doc.add_heading("1. Purpose", 1)
    doc.add_paragraph("This pack describes the capability model, the application catalogue and the strategic objectives of the Organization. It is used to test extraction.")
    for t in ("Capabilities are traceable to the mandate.", "Every application has an owner.", "Data is classified at source."):
        doc.add_paragraph(t, style="List Bullet")
    doc.add_paragraph("Review steps:")
    for t in ("Collect the evidence.", "Match it to a capability.", "Confirm with the owner."):
        doc.add_paragraph(t, style="List Number")

    doc.add_heading("2. Capability model", 1)
    doc.add_paragraph("The capability model has 6 level-1 groups. Colours show the current maturity of each capability.")
    doc.add_picture(str(out / "heat.png"), width=Inches(6.2))
    p = doc.add_paragraph("Figure 1: Capability map: supporting - enabling")
    p.runs[0].italic = True

    doc.add_heading("3. Application catalogue", 1)
    doc.add_heading("3.1 Evaluate", 2)
    t = doc.add_table(rows=len(APP_ROWS), cols=3)
    t.style = "Table Grid"
    for i, (label, value) in enumerate(APP_ROWS):
        a = t.cell(i, 0).merge(t.cell(i, 1))                      # the value spans two columns: Word stores it once, an extractor may repeat it
        a.text = value
        t.cell(i, 2).paragraphs[0].add_run(label).bold = True
    doc.add_heading("3.2 Attendance", 2)
    t = doc.add_table(rows=4, cols=3)
    t.style = "Table Grid"
    for i, (label, value) in enumerate([("Description", "Reads fingerprint machine data and passes the results to the HR & Payroll Portal by API or Excel."),
                                        ("Application Status", "In use"), ("Technologies Used", "Frontend: .NET DB: SQL Server"), ("Hosting", "On premises")]):
        t.cell(i, 0).merge(t.cell(i, 1)).text = value
        t.cell(i, 2).paragraphs[0].add_run(label).bold = True

    doc.add_heading("4. Capability catalogue", 1)
    doc.add_paragraph("Table 5: Capability catalogue, all level-1 and level-2 capabilities with their owner unit and maturity.")
    rows = 61
    big = doc.add_table(rows=rows, cols=6)
    big.style = "Table Grid"
    for j, h in enumerate(["Capability", "Level", "Group", "Owner unit", "Maturity", "Services"]):
        big.cell(0, j).text = h
        big.cell(0, j).paragraphs[0].runs[0].bold = True
        shade_header(big.cell(0, j))
    big.rows[0]._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))        # the header repeats on every page
    truth_rows = []
    groups = list(CAP_MAP)
    n = 0
    for gi, g in enumerate(groups):
        for cap, m in CAP_MAP[g]:
            n += 1
            row = [cap, "L2", g, f"OU-{10 + gi}", f"{m} " + ["Ad hoc", "Documented", "Evidenced", "Measured", "Optimized"][m - 1],
                   f"Services E{gi + 1}.{n % 5 + 1}: " + "supports " + cap.lower() + " for the whole organization and its branches " * 2]
            truth_rows.append(row)
    while len(truth_rows) < rows - 1:                                  # pad with level-1 rows so the table is long
        k = len(truth_rows)
        truth_rows.append([f"Level-1 capability {k:02d}", "L1", groups[k % len(groups)], f"OU-{20 + k % 9}", "2 Documented", f"Services C{k % 7 + 1}.{k % 4 + 1}"])
    truth_rows[41] = ["Job Performance", "L2", "Human Capital", "OU-11", "2 Documented", "Evaluation forms, bonus results and yearly performance review"]
    for i, row in enumerate(truth_rows[:rows - 1], 1):
        for j, v in enumerate(row):
            big.cell(i, j).text = v
    doc.add_page_break()

    doc.add_heading("5. Strategic objectives", 1)
    doc.add_paragraph("Table 10: Strategic objectives of the Organization")
    t = doc.add_table(rows=2 + len(OBJECTIVES), cols=3)
    t.style = "Table Grid"
    head = t.cell(0, 0).merge(t.cell(0, 2))
    head.text = "Strategic objectives 2026 to 2028"
    for j, h in enumerate(["Code", "Objective", "Intent"]):
        t.cell(1, j).text = h
        t.cell(1, j).paragraphs[0].runs[0].bold = True
    for i, row in enumerate(OBJECTIVES, 2):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    doc.add_paragraph()
    doc.add_paragraph("Table 11: Nested example, one objective with its KPIs")
    outer = doc.add_table(rows=2, cols=2)
    outer.style = "Table Grid"
    outer.cell(0, 0).text = "Objective"
    outer.cell(0, 1).text = "KPIs"
    outer.cell(1, 0).text = "SO4 Strengthen institutional excellence"
    inner = outer.cell(1, 1).add_table(rows=3, cols=2)
    for i, (k, v) in enumerate([("KPI", "Target"), ("Manual hand-offs in HR", "below 5 per month"), ("Evaluation to payroll lead time", "2 days")]):
        inner.cell(i, 0).text = k
        inner.cell(i, 1).text = v

    doc.add_heading("6. Architecture", 1)
    doc.add_paragraph("The evaluation results flow from the Evaluate Service to the HR & Payroll Portal. The portal feeds the Payroll Database and the Reporting Hub.")
    doc.add_picture(str(out / "arch.png"), width=Inches(6.0))
    p = doc.add_paragraph("Figure 2: Integration of the HR & Payroll Portal")
    p.runs[0].italic = True
    doc.add_picture(str(out / "chart.png"), width=Inches(5.6))
    p = doc.add_paragraph("Figure 3: Average capability maturity by level-1 capability")
    p.runs[0].italic = True
    box = parse_xml(
        f'<w:p {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml"><w:r><w:pict><v:shape style="width:300pt;height:50pt">'
        '<v:textbox><w:txbxContent><w:p><w:r><w:t>Callout: the digital strategy aims at a shared services capability for the whole organization.</w:t></w:r></w:p>'
        '</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>')
    doc.element.body.insert(len(doc.element.body) - 1, box)
    doc.add_paragraph("End of the pack. The mandate is shown in the digital strategy document.")
    doc.save(out / "Test_Architecture_Pack.docx")
    truth = {"cap_map": CAP_MAP, "arch": ARCH, "bars": BARS, "app_rows": APP_ROWS, "objectives": OBJECTIVES, "big_table_rows": truth_rows,
             "job_performance_row_index": 42}
    (out / "truth.json").write_text(json.dumps(truth, indent=1), encoding="utf-8")
    return out / "Test_Architecture_Pack.docx"


if __name__ == "__main__":
    print(build(Path(sys.argv[1] if len(sys.argv) > 1 else "extract_test")))
