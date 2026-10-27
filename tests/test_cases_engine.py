"""A sheet of demands runs through Demand Management: gate, properties in order, questions, approval."""
import json
import re
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from eawiki import history  # noqa: E402

history.DB_PATH = Path(tempfile.mkdtemp()) / "cases_test.db"

from eawiki.cases import gather as gmod  # noqa: E402
from eawiki.cases import intake, plan as pmod, readiness, summary  # noqa: E402
from eawiki.cases.engine import Engine, levels  # noqa: E402
from eawiki.cases.store import CaseStore  # noqa: E402

SPEC = json.loads((ROOT / "use_cases" / "demand-management.json").read_text(encoding="utf-8"))
ITEMS = intake.parse_file(SPEC, "x.csv", (ROOT / "use_cases" / "samples" / "demand-management-inputs.csv").read_bytes())["items"]
ITEMS.append({"demand_id": "D-2026-006", "title": "VPN not working for remote staff", "business_need": "VPN drops every 20 minutes",
              "expected_outcome": "Stable VPN", "department": "IT", "requester_role": "Analyst", "sponsor": "CIO",
              "date_submitted": "2026-08-20", "request_type": "New solution", "requester_urgency": "High",
              "target_go_live": "2026-09-01"})

CANNED = {
    "Demand type": "Digital", "Domain": ["Application", "Security"], "Urgency": "High",
    "Priority": "Standard", "Mandate / business or DT objective": "O-02 / I-04",
    "Linked initiative or project": None,                           # cites nothing: the architect is asked
    "Impacted business capability": ["Employee Self-Service"], "Impacted application": ["HR system (C-204)"],
    "Is duplicate": "No", "Duplicate demand number (in the sheet)": "", "Duplicate with existing artefact / artefact name": "None",
    "Architecture principles and standards alignment": "Conforms", "Data classification": "Confidential",
    "Linked security capability": ["Identity and access management"], "Security review needed": "Yes",
    "Screening observation": "A digital demand.",
    "Recommendation": "Enhance existing", "Overlaps with": ["HR system"], "Consolidate with demand(s)": [],
    "Contract / licence match": "None", "Security fit": "Fits", "Recommended owner": "HR Systems Lead",
    "Justification": "Overlap with the HR system.", "Confidence": "Medium",
}


def fake_llm(prompt: str):
    if "plan WHAT TO LOOK FOR" in prompt:
        return {"sources": ["demand", "documents"], "documents": [], "bogus": 1,
                "lookups": [{"find": "what each system named is", "from": "systems", "match": "name", "in": ["Strategy"], "extract": "its owner"},
                            {"find": "invalid", "from": "bogus", "match": "name"}]}
    if "list what it MENTIONS" in prompt:
        m = re.search(r"Systems or data mentioned: (.+)", prompt)
        return {"systems": [x.strip() for x in m.group(1).split(",")] if m else [], "data": ["a thing nobody wrote"]}
    if "INTAKE QUESTIONS" in prompt:
        return {"outcome": "No: incident", "reason": "The VPN worked before."} if "VPN" in prompt else {"outcome": "Yes", "reason": ""}
    m = re.search(r"PROPERTY: (.+?)  \[", prompt)
    name = m.group(1)
    if "THE ARCHITECT POINTED YOU HERE" in prompt and "document Roadmap" in prompt:
        ids = re.findall(r"\[([^\]\s]+:\d+)\]", prompt)
        return {"value": "P-12 HR cloud move", "status": "answered", "evidence": ids[:1], "needs": "", "note": "found in the roadmap"}
    val = CANNED.get(name)
    ids = re.findall(r"\[([^\]\s]+:\d+)\]", prompt)
    if val is None:
        return {"value": "I Guess", "status": "answered", "evidence": [], "needs": "", "note": "guess",
                "question": "Which roadmap project does this demand belong to?",
                "suggestions": [{"value": "P-103 HR cloud move", "why": "the roadmap lists it", "evidence": ids[:1]},
                                {"value": "P-999 invented", "why": "no source", "evidence": ["Nope:9"]},
                                {"value": "P-104 no evidence", "why": "unsupported", "evidence": []}]}
    ev = ids[:1] if name == "Mandate / business or DT objective" else []
    return {"value": val, "status": "answered", "evidence": ev, "needs": "", "note": "from the evidence"}


def fake_tools():
    block = {"block_id": "Strategy:1", "doc_id": "Strategy", "text": "O-02 employee requests handled digitally."}
    road = {"block_id": "Roadmap:7", "doc_id": "Roadmap", "text": "P-12 HR cloud move, owner HR."}
    return gmod.Tools(search=lambda q, k, docs: [road] if docs == ["Roadmap"] else ([] if docs == ["Gap_Analysis_Sample"] else [block]), blocks_by_ids=lambda ids: [block],
                      link=lambda t: {"entities": [{"mention": "HR system", "ids": ["hr"], "names": ["HR system"]}]},
                      graph_facts=lambda ids: {"facts": "HR system --is_deployed_on--> Oracle DB", "block_ids": ["Strategy:1"]},
                      llm_json=fake_llm, doc_ids=lambda: ["Strategy", "Gap_Analysis_Sample", "Roadmap"])


def _engine():
    return Engine(CaseStore(), fake_tools, workers=4)


def test_levels_put_a_property_after_what_it_builds_on():
    cols = SPEC["stages"][0]["output"]["columns"]
    order = [[c["name"] for c in lv] for lv in levels([c for c in cols if c["name"] not in ("Screening observation", "Evidence")])]
    flat = [n for lv in order for n in lv]
    assert flat.index("Urgency") > flat.index("Mandate / business or DT objective")
    assert flat.index("Priority") > flat.index("Urgency")


def test_a_plan_cannot_point_at_something_that_does_not_exist():
    col = {"name": "X", "look_in": ["Strategy"]}
    p = pmod.validate({"sources": ["demand", "nonsense", "documents"], "documents": ["strategy", "Ghost"],
                       "lookups": [{"find": "a system", "from": "systems", "match": "name", "in": ["Application Catalog", "Ghost"]},
                                   {"find": "bad", "from": "nonsense", "match": "name"},
                                   {"find": "", "from": "systems", "match": "name"}]},
                      col, ["business_need"], ["Strategy", "Application_Catalog_v2", "Roadmap"])
    assert p["sources"] == ["demand", "documents"] and p["documents"] == ["Strategy", "Application_Catalog_v2"]
    assert len(p["lookups"]) == 1 and p["lookups"][0]["docs"] == ["Application_Catalog_v2"] and p["lookups"][0]["in"] == ["Application Catalog", "Ghost"]
    # a kind of source finds a document added later, because it is matched by its words, not listed by name
    assert pmod.resolve_docs("Business Capability Map", ["Business_Capability_Map_2027", "Strategy"]) == ["Business_Capability_Map_2027"]
    # lookups written by the owner are used as they are: no model call
    col2 = {"name": "Y", "lookups": [{"find": "f", "from": "need", "match": "meaning", "in": ["Roadmap"]}]}
    q = pmod.build_plan(col2, ["business_need"], ["Roadmap"], llm=lambda _p: (_ for _ in ()).throw(AssertionError("no call")))
    assert q["origin"] == "config" and q["lookups"][0]["docs"] == ["Roadmap"]


def test_the_profile_keeps_only_what_the_requester_wrote():
    from eawiki.cases import profile
    f = {"systems_mentioned": ("Systems or data mentioned", "HR system, payroll, attendance devices"),
         "business_need": ("Business need", "Employees request leave by email.")}
    p = profile.validate({"systems": ["HR system", "payroll", "the SAP finance suite"], "data": ["a thing nobody wrote"],
                          "processes": ["leave"]}, f)
    assert [m["text"] for m in p["systems"]] == ["HR system", "payroll"] and p["data"] == []
    assert p["systems"][0]["field"] == "systems_mentioned" and p["processes"][0]["field"] == "business_need"
    assert profile.build_profile(f, None)["origin"] == "none"


def test_lookups_search_each_mention_and_say_when_nothing_names_it():
    tools = fake_tools()
    plan = {"sources": ["demand", "documents"], "queries": ["{__text__}"], "documents": ["Strategy"], "origin": "t",
            "lookups": [{"find": "what each system is", "from": "systems", "match": "name", "in": ["Strategy"], "docs": ["Strategy"], "extract": ""}]}
    prof = {"systems": [{"text": "digitally", "field": "x"}, {"text": "zeppelin", "field": "x"}], "origin": "model"}
    item = dict(ITEMS[0], _id="D-2026-001")
    ev = gmod.gather(tools, plan, SPEC, item, [], {}, None, prof)
    rec = ev["lookups"][0]
    assert [i["mention"] for i in rec["items"]] == ["digitally", "zeppelin"]
    assert rec["items"][0]["none"] is False and rec["items"][1]["none"] is True       # one named in a passage, one not
    txt = gmod.render(ev)
    assert 'NO passage names it' in txt and "has an entry" in txt and "LOOKUPS" in txt and any("no match" in x for x in ev["searched"])
    # a passage that only mentions a name is not an entry of its own
    plan2 = dict(plan, lookups=[dict(plan["lookups"][0], docs=["Strategy"])])
    t2 = fake_tools()
    t2.search = lambda q, k, docs: [{"block_id": "Strategy:9", "doc_id": "Strategy", "text": "C-204 Peoplex HR system. Interfaces to a payroll export. Owner: HR Systems Lead."}]
    ev3 = gmod.gather(t2, plan2, SPEC, item, [], {}, None, {"systems": [{"text": "payroll", "field": "x"}, {"text": "Peoplex HR", "field": "x"}], "origin": "model"})
    a, b = ev3["lookups"][0]["items"]
    assert a["none"] is True and a["mentioned"] == ["Strategy:9"] and b["none"] is False
    assert "NO entry of its own" in gmod.render(ev3) and any("only mentioned in" in x for x in ev3["searched"])
    # a profile that could not be built falls back to the whole text, and never claims "mentions no systems"
    ev2 = gmod.gather(tools, plan, SPEC, item, [], {}, None, {"origin": "none"})
    assert ev2["lookups"][0]["items"][0]["mention"] == "the demand as a whole"


def test_a_sheet_runs_asks_for_what_is_missing_and_moves_to_the_next_phase():
    e = _engine()
    cid = e.create(SPEC, ITEMS, "test")
    e.run_phase(cid, "screening")
    st = e.store
    cells = {(c["item_id"], c["column"]): c for c in st.cells(cid, "screening")}
    n_cols = len(SPEC["stages"][0]["output"]["columns"])
    assert len({c["item_id"] for c in cells.values()}) == 4 and len(cells) == 4 * n_cols
    assert cells[("D-2026-001", "Is it a true demand?")]["value"] == "Yes"
    # the incident is stopped at the gate: every other property reads Not assessed, with the reason
    assert cells[("D-2026-006", "Is it a true demand?")]["value"] == "No: incident"
    assert cells[("D-2026-006", "Domain")]["status"] == "not_assessed"
    assert cells[("D-2026-006", "Priority")]["value"] == "No score"
    assert "No: incident" in cells[("D-2026-006", "Screening observation")]["value"]
    # a property that needs a source and cited none is not answered: it opens a question
    c = cells[("D-2026-001", "Linked initiative or project")]
    assert c["status"] == "needs_input" and c["value"] == "" and "roadmap" in c["need"]
    qs = [q for q in st.questions(cid) if q["status"] == "open"]
    assert {q["column"] for q in qs} >= {"Linked initiative or project"} and len(qs) >= 3
    # the question is the model's specific one; its suggestions kept only when a shown source supports them;
    # what was searched is recorded from the results
    q0 = next(q for q in qs if q["item_id"] == "D-2026-001" and q["column"] == "Linked initiative or project")
    assert q0["text"] == "Which roadmap project does this demand belong to?"
    assert [x["value"] for x in q0["meta"]["suggestions"]] == ["P-103 HR cloud move"] and q0["meta"]["suggestions"][0]["evidence"] == ["Strategy:1"]
    assert "passage" in q0["meta"]["looked"] and "Strategy" in q0["meta"]["looked"]
    # a cited source is real; a property that needs none may rest on the demand
    assert cells[("D-2026-001", "Mandate / business or DT objective")]["evidence"] == ["Strategy:1"]
    assert cells[("D-2026-001", "Priority")]["status"] == "answered"
    # the citations property is the union, in code
    assert cells[("D-2026-001", "Evidence")]["value"] == ["Strategy:1"]
    # the summary is computed from the cells
    sm = summary.phase_summary(SPEC, SPEC["stages"][0], st.cells(cid), e.store.get_case(cid)["items"], st.questions(cid))
    assert sm["assessed"] == 3 and sm["total"] == 4 and sm["returned"][0]["id"] == "D-2026-006"
    assert {k["label"]: k["value"] for k in sm["kpis"]}["Security review"] == 3
    assert "Which demands were returned, and why?" in sm["suggested"]
    # the architect answers: stored as theirs, and what builds on it is decided again
    q = next(q for q in qs if q["item_id"] == "D-2026-001" and q["column"] == "Linked initiative or project")
    out = e.answer(q["id"], "P-103 HR system cloud migration", background=False, via="suggestion")
    assert out["status"] == "answered"
    assert e.store.get_question(q["id"])["answer"].startswith("suggestion: ")
    again = {c["column"]: c for c in st.cells(cid, "screening") if c["item_id"] == "D-2026-001"}
    assert again["Linked initiative or project"]["status"] == "user" and again["Linked initiative or project"]["source"] == "user"
    # approve: the next phase runs for the assessed demands only
    assert st.get_case(cid)["state"]["phases"]["screening"]["status"] == "review"
    e.approve(cid, "screening", background=False)
    rec = {(c["item_id"], c["column"]) for c in st.cells(cid, "recommendation")}
    assert {i for i, _ in rec} == {"D-2026-001", "D-2026-003", "D-2026-004"}


def test_the_architect_can_point_at_a_document_instead_of_giving_the_value():
    e = _engine()
    cid = e.create(SPEC, ITEMS[:1], "look")
    e.run_phase(cid, "screening")
    q = next(q for q in e.store.questions(cid) if q["status"] == "open" and q["column"] == "Linked initiative or project")
    for bad in (dict(document="Nope", where=""), dict(document=None, where="")):
        try:
            e.look(q["id"], bad["document"], bad["where"], background=False)
            assert False, "should refuse"
        except ValueError as err:
            assert str(err)
    # a document that has nothing: the question stays open and says what was tried; nothing is guessed
    e.look(q["id"], "Strategy", "the roadmap row", background=False)
    q2 = e.store.get_question(q["id"])
    assert q2["status"] == "open" and "document Strategy" in q2["meta"]["looked"] and "Nothing there decides" in q2["meta"]["looked"]
    # the right document: answered with a real citation, the question closes
    e.look(q["id"], "Roadmap", "HR cloud move", background=False)
    assert e.store.get_question(q["id"])["status"] == "answered"
    c = next(c for c in e.store.cells(cid, "screening") if c["column"] == "Linked initiative or project")
    assert c["status"] == "answered" and c["source"] == "model" and c["evidence"] == ["Roadmap:7"]
    assert c["how"][0].startswith("The architect pointed to: document Roadmap")


def test_the_flows_are_langgraph_graphs_that_can_be_drawn():
    from eawiki.cases import flow
    d = flow.draw(_engine())
    for node in ("prepare", "gate", "run_round", "summaries", "citations", "finish"):
        assert node in d["phase"]
    for node in ("gather", "solve", "save_answer", "open_question", "look_found", "look_missed"):
        assert node in d["cell"]


def test_duplicates_name_each_other_and_a_disagreement_is_asked():
    e = _engine()
    items = [dict(ITEMS[0], _id="D-2026-001"), dict(ITEMS[1], _id="D-2026-003"), dict(ITEMS[2], _id="D-2026-004")]
    cid = e.create(SPEC, items, "dups")
    st = e.store
    F, R = "Is duplicate", "Duplicate demand number (in the sheet)"
    for iid, flag, ref in (("D-2026-001", "Yes", "D-2026-003; D-2026-004"), ("D-2026-003", "Yes", "None"), ("D-2026-004", "No", "None")):
        st.save_cell(cid, iid, "screening", F, value=flag, status="answered")
        st.save_cell(cid, iid, "screening", R, value=ref, status="answered", how=["from the model"])
    stage = SPEC["stages"][0]
    e._n_consistency({"cid": cid, "pid": "screening", "only_item": None, "run": {"stage": stage, "rows": items}})
    cells = {(c["item_id"], c["column"]): c for c in st.cells(cid, "screening")}
    assert cells[("D-2026-003", R)]["value"] == "D-2026-001" and cells[("D-2026-003", R)]["source"] == "rule"
    assert "mirrored in code" in cells[("D-2026-003", R)]["how"][-1]
    q = [x for x in st.questions(cid) if x["status"] == "open"]
    assert len(q) == 1 and q[0]["item_id"] == "D-2026-004" and q[0]["column"] == F        # one says Yes, the other No: a person decides


def test_a_case_is_scored_against_the_expected_answers():
    from eawiki.cases import score as sc
    gold = json.loads((ROOT / "use_cases" / "samples" / "demand-management-gold.json").read_text(encoding="utf-8"))
    assert len(gold["items"]) == 9 and "Domain" in gold["items"]["D-2026-001"]["expected"]
    assert sc.compare("O-04 / I-08; O-03", "O-04 and O-03, I-08 (the mandate)") == "match"
    assert sc.compare("Application; Security", "security; application") == "match"
    assert sc.compare("Application; Security", "Application") == "partial"
    assert sc.compare("Standard", "Fast-track") == "wrong"
    cells = [{"phase_id": "screening", "item_id": "D-2026-001", "column": "Domain", "value": ["Application", "Security"], "status": "answered"},
             {"phase_id": "screening", "item_id": "D-2026-001", "column": "Priority", "value": "Fast-track", "status": "answered"},
             {"phase_id": "screening", "item_id": "D-2026-001", "column": "Data classification", "value": "", "status": "needs_input"}]
    rep = sc.score(cells, gold)
    assert rep["properties"]["Priority"]["wrong"] == 1 and rep["properties"]["Data classification"]["asked"] == 1
    assert "ALL" in sc.render(rep) and any(m["property"] == "Priority" for m in rep["misses"])


def test_the_verified_sheet_is_loaded_into_the_use_case_by_property_name():
    from openpyxl import Workbook
    from eawiki import usecases
    from eawiki.cases import import_sheet as imp
    wb = Workbook()
    ws = wb.active
    ws.title = "Properties corrected"
    ws.append(["#", "Field", "Type", "Values (exact words)", "Builds on (decided first)", "Source required?", "Min sources",
               "Description: what this field means", "How to resolve it (steps)", "Where to look (x)", "Rules / decision criteria",
               "Examples (one correct, one wrong)", "Info needed: when INPUT is missing (ask)", "Info needed: when a SOURCE is missing (ask)"])
    ws.append(["D1", "Stage instructions: Screening", "Text", "", "", "n/a", "", "d", "h", "w", "Never approve or reject.", "e", "i", "s"])
    ws.append(["D2", "Domain", "List", "Business; Data; Application; Technology; Security", "Data classification", "No (the demand can decide)", "1 when a catalog is used",
               "The architecture domains.", "1. Read.", "Application Catalog", "Five domains only.", "Correct: Application\n[PROPOSED by Claude: replace]",
               "Ask which systems.", "[PROPOSED by Claude] Say so."])
    ws.append(["D4", "Priority", "Choice", "Fast-track; Standard; Backlog", "Mandate / objective; Is duplicate", "Yes", "2 (a judgment)",
               "d", "h", "w", "r", "e", "i", "s"])
    ws.append(["D3", "Category", "List", "", "", "No", "", "d", "h", "w", "r", "e", "i", "s"])
    ws.append(["D5", "Mandate / objective", "Text", "", "", "Yes", "", "[EA to write]", "", "", "", "", "", ""])
    p = ROOT / "tests" / "_sheet_tmp.xlsx"
    wb.save(p)
    try:
        new, rep = imp.apply(SPEC, str(p))
    finally:
        p.unlink()
    dom = next(c for c in new["stages"][0]["output"]["columns"] if c["name"] == "Domain")
    assert dom["values"] == ["Business", "Data", "Application", "Technology", "Security"] and dom["uses"] == ["Data classification"]
    assert dom["evidence_required"] is False and dom["examples"] == "Correct: Application" and dom["info_needed_source"] == "Say so."
    pri = next(c for c in new["stages"][0]["output"]["columns"] if c["name"] == "Priority")
    assert pri["min_sources"] == 2 and pri["uses"] == ["Mandate / business or DT objective", "Is duplicate"]
    assert new["playbook"] == "Never approve or reject." and rep["unmatched"] == ["Category"]
    assert "Mandate / business or DT objective" in rep["skipped"]                  # only a note to the reader: left alone
    assert usecases.validate(new) == []
    assert "examples" not in next(c for c in SPEC["stages"][0]["output"]["columns"] if c["name"] == "Domain")      # the original is untouched


def test_a_list_property_keeps_to_its_pick_list_and_gets_one_retry():
    from eawiki.cases import solve as smod
    col = {"name": "Domain", "type": "list", "values": ["Business", "Data", "Application", "Technology", "Security"],
           "evidence_required": False}
    ev = {"demand": {}, "blocks": [], "graph": None, "others": [], "allowed": {"demand:x"}, "used": [], "searched": []}
    calls = []

    def llm(prompt):
        calls.append(prompt)
        if "REJECTED" in prompt:
            return {"value": ["Application", "security"], "status": "answered", "evidence": ["demand:x"]}
        return {"value": ["Application", "Integration"], "status": "answered", "evidence": ["demand:x"]}
    cell = smod.solve(col, {}, SPEC, ev, {}, llm)
    assert cell["status"] == "answered" and cell["value"] == ["Application", "Security"] and len(calls) == 2
    assert "'Integration' is not one of" in calls[1] and "second was accepted" in cell["how"][-1]
    stubborn = smod.solve(col, {}, SPEC, ev, {}, lambda p: {"value": ["Integration"], "status": "answered", "evidence": ["demand:x"]})
    assert stubborn["status"] == "needs_input" and "not one of" in stubborn["need"]


def test_the_use_case_names_its_items_and_what_each_property_is_doing():
    from eawiki import usecases
    assert SPEC["input"]["item_label"] == {"singular": "demand", "plural": "demands"}
    assert usecases.validate(SPEC) == []
    bad = json.loads(json.dumps(SPEC))
    bad["input"]["item_label"] = {"singular": "demand"}
    bad["stages"][0]["output"]["columns"][1]["activity"] = "x" * 200
    errs = usecases.validate(bad)
    assert any("item label" in x for x in errs) and any("activity" in x for x in errs)
    st = [x["label"] for x in __import__("eawiki.cases.engine", fromlist=["steps"]).steps(SPEC, SPEC["stages"][0], True)]
    assert st[0] == "Checking each demand is a true demand" and "Deciding the demand type" in st


def test_the_phase_shows_as_running_before_the_plans_are_built_and_a_failure_is_recorded():
    seen = []
    e = _engine()
    cid = e.create(SPEC, ITEMS[:1], "early")
    base = fake_tools

    def tools():
        t = base()
        inner = t.llm_json

        def llm(prompt):
            if "plan WHAT TO LOOK FOR" in prompt:           # a plan is being built
                seen.append(e.store.get_case(cid)["state"]["phases"].get("screening", {}).get("status"))
            return inner(prompt)
        t.llm_json = llm
        return t
    e.tools_factory = tools
    e.run_phase(cid, "screening")
    assert seen and set(seen) == {"running"}                                   # the screen already shows the run

    def broken():
        raise RuntimeError("the model service is down")
    e2 = _engine()
    cid2 = e2.create(SPEC, ITEMS[:1], "broken")
    e2.tools_factory = broken
    try:
        e2.run_phase(cid2, "screening")
        assert False, "should raise"
    except RuntimeError:
        pass
    ph = e2.store.get_case(cid2)["state"]["phases"]["screening"]
    assert ph["status"] == "error" and "model service is down" in ph["error"]
    assert any(ev["kind"] == "error" for ev in e2.store.events_after(cid2, 0))


def test_resume_keeps_what_was_decided_and_the_architects_answers():
    calls = []
    e = _engine()
    cid = e.create(SPEC, ITEMS[:1], "resume")
    base = fake_tools

    def tools():
        t = base()
        inner = t.llm_json

        def llm(prompt):
            if "PROPERTY:" in prompt:
                calls.append(re.search(r"PROPERTY: (.+?)  \[", prompt).group(1))
            return inner(prompt)
        t.llm_json = llm
        return t
    e.tools_factory = tools
    e.run_phase(cid, "screening")
    q = next(q for q in e.store.questions(cid) if q["column"] == "Linked initiative or project")
    e.answer(q["id"], "P-103 HR cloud move", background=False)
    first = len(calls)
    calls.clear()
    e.resume(cid, "screening", background=False)
    assert "Domain" not in calls and "Priority" not in calls                  # decided: not asked again
    assert "Linked initiative or project" not in calls                       # the architect's answer is kept
    c = next(c for c in e.store.cells(cid, "screening") if c["column"] == "Linked initiative or project")
    assert c["status"] == "user" and c["value"] == "P-103 HR cloud move"
    assert len(calls) < first


def test_readiness_names_the_missing_document_and_the_empty_source():
    spec = json.loads(json.dumps(SPEC))
    spec["stages"][0]["output"]["columns"][5]["look_in"] = ["Roadmap_Not_Ingested"]
    rep = readiness.readiness(spec, [dict(i, _id=i["demand_id"]) for i in ITEMS[:3]], fake_tools())
    props = {p["property"]: p for p in rep["phases"][0]["properties"]}
    p = props["Urgency"]
    assert p["verdict"] in ("ready", "partial")
    miss = props[spec["stages"][0]["output"]["columns"][5]["name"]]
    assert miss["verdict"] == "not_ready" and "not ingested" in miss["reasons"][0]


def test_api_runs_in_the_background_and_the_screen_can_follow_it():
    from fastapi.testclient import TestClient
    from eawiki.api import cases as api
    from eawiki.api.app import app
    api.set_engine(_engine())
    client = TestClient(app)
    r = client.post("/api/cases/parse-file", params={"use_case_id": "demand-management"},
                    files={"file": ("d.csv", (ROOT / "use_cases" / "samples" / "demand-management-inputs.csv").read_bytes())})
    assert r.status_code == 200 and len(r.json()["items"]) == 3
    cid = client.post("/api/cases", json={"use_case_id": "demand-management", "items": r.json()["items"], "name": "t"}).json()["id"]
    for _ in range(100):
        d = client.get(f"/api/cases/{cid}").json()
        if d["phases"][0]["state"].get("status") == "review":
            break
        time.sleep(0.1)
    assert d["phases"][0]["state"]["status"] == "review"
    assert d["cells"]["screening"]["D-2026-001"]["Domain"]["display"] == "answered"
    assert any(e["kind"] == "phase_done" for e in d["events"]) and d["last_seq"] >= 1
    x = client.get(f"/api/cases/{cid}/export", params={"phase": "screening"})
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"]
    assert d["item_label"]["plural"] == "demands" and d["phases"][0]["steps"][0]["label"] == "Checking each demand is a true demand"
    qd = next(q for q in d["questions"] if q["column"] == "Linked initiative or project")
    assert qd["options"] == [] and qd["meta"]["suggestions"]
    import io
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(x.content)).active
    assert [c.value for c in ws[1]][:3] == ["Demand", "Title", "Is it a true demand?"] and ws.max_row == 4
    assert client.patch(f"/api/cases/{cid}", json={"name": "  Renamed   case "}).json()["name"] == "Renamed case"
    assert client.patch(f"/api/cases/{cid}", json={"name": "   "}).status_code == 400
    assert client.get(f"/api/cases/{cid}").json()["case"]["name"] == "Renamed case"
    assert client.post(f"/api/cases/{cid}/phases/recommendation/approve").status_code == 409   # not ready for review
    assert client.delete(f"/api/cases/{cid}").json()["deleted"] is True
    assert client.get(f"/api/cases/{cid}").status_code == 404
    assert client.patch(f"/api/cases/{cid}", json={"name": "x"}).status_code == 404
    api.set_engine(None)


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
    print("ok")
