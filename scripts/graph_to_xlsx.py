"""Export the full extracted graph (data/graph/graph.json) to a shareable Excel workbook.

This is the PIPELINE OUTPUT — every entity and relationship extracted across all docs —
for the EA team to read/review. (Different from eval/gold.xlsx, which is the small
hand-labelled grading key.)

    python scripts/graph_to_xlsx.py            # -> exports/architecture-graph.xlsx
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
HEAD_FILL = PatternFill("solid", fgColor="2D6CC0")
HEAD_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")


def _sheet(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c, w in zip(ws[1], widths):
        c.fill = HEAD_FILL; c.font = HEAD_FONT
        ws.column_dimensions[c.column_letter].width = w
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    ws.freeze_panes = "A2"


def main() -> None:
    graph = json.loads((ROOT / "data/graph/graph.json").read_text(encoding="utf-8"))
    menu = json.loads((ROOT / "data/togaf92_metamodel.json").read_text(encoding="utf-8"))["node_types"]
    comms = {}
    cpath = ROOT / "data/graph/communities.json"
    if cpath.exists():
        for c in json.loads(cpath.read_text(encoding="utf-8")).get("communities", []):
            for mid in c.get("member_ids", []):
                comms[mid] = c["name"]

    nodes = graph["nodes"]
    id2name = {n["node_id"]: n.get("name", n["node_id"]) for n in nodes}
    domain = lambda t: (menu.get(t, {}) or {}).get("domain", "Other")

    ent_rows = [[
        n.get("name", ""), n.get("type", ""), domain(n.get("type", "")),
        comms.get(n["node_id"], ""), "; ".join(n.get("kinds", [])),
        "; ".join(n.get("aliases", [])), "; ".join(n.get("origins", [])),
        len(n.get("block_ids", [])),
    ] for n in sorted(nodes, key=lambda n: (domain(n.get("type", "")), n.get("name", "")))]

    rel_rows = [[
        id2name.get(e["source"], e["source"]), e.get("relation", ""),
        id2name.get(e["target"], e["target"]), e.get("source_count", 1),
        "; ".join(e.get("kinds", [])), "; ".join(e.get("origins", [])),
    ] for e in sorted(graph["edges"], key=lambda e: id2name.get(e["source"], ""))]

    wb = Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("README"); ws.column_dimensions["A"].width = 100
    ws["A1"] = "Enterprise Architecture — Extracted Knowledge Graph"
    ws["A1"].font = Font(bold=True, size=14)
    for i, ln in enumerate([
        "",
        "The full graph extracted from the EA documents by the pipeline (entities + typed",
        "relationships), grounded in the sources. Types follow the TOGAF 9.2 content metamodel.",
        "",
        f"Entities: {len(nodes)}    Relationships: {len(graph['edges'])}    Communities: {len(set(comms.values()))}",
        "",
        "Columns — Relationships: 'Sources' = which extraction pass(es) found it; 'Corroboration' =",
        "how many independent sources asserted it (higher = more confident).",
        "This is the extracted output for review — not a hand-verified gold set.",
    ], start=2):
        ws[f"A{i}"] = ln; ws[f"A{i}"].alignment = WRAP

    _sheet(wb, "Entities", ["Name", "Type", "Domain", "Community", "Doc kinds", "Aliases", "Sources", "# source blocks"],
           ent_rows, [34, 30, 14, 24, 16, 28, 20, 14])
    _sheet(wb, "Relationships", ["Source", "Relation", "Target", "Corroboration", "Doc kinds", "Sources"],
           rel_rows, [32, 20, 32, 14, 16, 22])

    out = ROOT / "exports/architecture-graph.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    print(f"wrote {out}  ({len(nodes)} entities, {len(graph['edges'])} relationships)")


if __name__ == "__main__":
    main()
